from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def create_assessment_output(assessment_history: list) -> BytesIO:
    """
    Creates an Excel workbook containing all assessment results.

    Each assessment occupies one row.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "Objection Assessments"

    # -----------------------------
    # Header
    # -----------------------------
    headers = [
        "Objection",
        "Complexity",
        "Triggered Rubric Criteria",
        "Relevant PTA Sections",
        "Recommendation",
        "Similar Past Cases",
    ]

    ws.append(headers)

    # Bold header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # -----------------------------
    # Assessment rows
    # -----------------------------
    for assessment in assessment_history:

        grounds = assessment["grounds_of_objection"]

        complexity = assessment["classification"]["complexity"]

        rubric = "\n".join(
            assessment["classification"]["rubric_criteria"]
        )

        pta_sections = "\n\n".join(
            f"{section['section']}\n{section['content']}"
            for section in assessment["classification"]["relevant_pta_sections"]
        )

        recommendation = "\n".join(
                    assessment["recommendation"]["next_steps"]
                )

        similar_cases = ""

        for i, case in enumerate(
            assessment["similar_cases"],
            start=1,
        ):
            similar_cases += (
                f"Case {i}\n"
                f"Property Type: {case['property_type']}\n"
                f"Development: {case['development']}\n"
                f"Strata Classification: {case['strata_classification']}\n"
                f"Objection: {case['grounds_of_objection']}\n"
                f"File Upload Count: {case['file_upload_count']}\n"
                f"Complexity: {case['complexity']}\n"
                f"Reasoning: {case['reasoning']}\n\n"
            )

        ws.append([
            grounds,
            complexity,
            rubric,
            pta_sections,
            recommendation,
            similar_cases.strip(),
        ])

    # -----------------------------
    # Formatting
    # -----------------------------
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    column_widths = {
        "A": 45,
        "B": 20,
        "C": 35,
        "D": 45,
        "E": 50,
        "F": 60,
    }

    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # -----------------------------
    # Return workbook as bytes
    # -----------------------------
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output